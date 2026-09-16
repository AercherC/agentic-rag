"""Phase 14A-A 受控语料生成器（子进程 runner，运行于隔离沙盒）。

读取 corpus14a.DOCX_CASES / XLSX_CASES（纯 Python 意图），写出真实 .docx / .xlsx。
Gold 由主 harness 从同一 spec 写出（selfcheck 验证"写出 == 意图"），候选 parser 不参与 Gold。

运行方式（由 adapters14a 子进程调用）：
    python -m phase14a.runners.build_corpus_arm --out <dir>
其中 python = 隔离沙盒解释器（PYTHONPATH 含 experiments + 隔离 site-packages）。
"""

from __future__ import annotations

import sys
import io
import base64
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from phase14a import corpus14a, sandbox14a

# 1x1 透明 PNG，用于 image/caption 用例（避免外部图片依赖）。
# 由 zlib 现生成的合法 RGBA PNG，避免手敲 base64 截断导致 docx 解析失败。
_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAAC0lEQVR4nGNgAAIAAAUAAXpeqz8"
    "AAAAASUVORK5CYII="
)


def _png_stream() -> io.BytesIO:
    return io.BytesIO(base64.b64decode(_PNG_B64))


def _xml_escape(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _add_textbox(document: Document, text: str) -> None:
    """尝试插入真实浮动 textbox；失败则退化为单格表格 callout（不阻塞生成）。"""

    try:
        para = document.add_paragraph()
        run = para.add_run()
        d = OxmlElement("w:drawing")
        anchor = OxmlElement("wp:anchor")
        for k, v in [
            ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
            ("simplePos", "0"), ("relativeHeight", "251658240"), ("behindDoc", "0"),
            ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1"),
        ]:
            anchor.set(qn("wp:" + k), v)
        sp = OxmlElement("wp:simplePos"); sp.set(qn("wp:x"), "0"); sp.set(qn("wp:y"), "0")
        ph = OxmlElement("wp:positionH"); ph.set(qn("wp:relativeFrom"), "column")
        ph_o = OxmlElement("wp:posOffset"); ph_o.text = "0"
        pv = OxmlElement("wp:positionV"); pv.set(qn("wp:relativeFrom"), "paragraph")
        pv_o = OxmlElement("wp:posOffset"); pv_o.text = "0"
        ext = OxmlElement("wp:extent"); ext.set(qn("wp:cx"), "2400000"); ext.set(qn("wp:cy"), "1200000")
        eff = OxmlElement("wp:effectExtent")
        for k, v in [("l", "0"), ("t", "0"), ("r", "0"), ("b", "0")]:
            eff.set(qn("wp:" + k), v)
        docpr = OxmlElement("wp:docPr"); docpr.set(qn("wp:id"), "1"); docpr.set(qn("wp:name"), "TextBox")
        wsp = OxmlElement("wps:wsp")
        cnv = OxmlElement("wps:cnvPr"); cnv.set(qn("wps:id"), "1"); cnv.set(qn("wps:name"), "TextBox")
        sppr = OxmlElement("wps:spPr")
        xfrm = OxmlElement("a:xfrm")
        off = OxmlElement("a:off"); off.set("x", "0"); off.set("y", "0")
        aext = OxmlElement("a:ext"); aext.set("cx", "2400000"); aext.set("cy", "1200000")
        xfrm.append(off); xfrm.append(aext)
        geom = OxmlElement("a:prstGeom"); geom.set("prst", "rect")
        sppr.append(xfrm); sppr.append(geom)
        txbx = OxmlElement("wps:txbx")
        content = OxmlElement("w:txbxContent")
        tp = OxmlElement("w:p"); tr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = text
        tr.append(tt); tp.append(tr); content.append(tp); txbx.append(content)
        wsp.append(cnv); wsp.append(sppr); wsp.append(txbx)
        bpr = OxmlElement("wps:bodyPr")
        anchor.append(sp); anchor.append(ph); anchor.append(ph_o); anchor.append(pv)
        anchor.append(pv_o); anchor.append(ext); anchor.append(eff); anchor.append(docpr)
        anchor.append(wsp); anchor.append(bpr); d.append(anchor); run._r.append(d)
    except Exception:  # 退化：单格表格 callout
        tbl = document.add_table(rows=1, cols=1)
        tbl.cell(0, 0).text = f"[TEXTBOX] {text}"
        tbl._cells_fallback = True  # type: ignore[attr-defined]


def _build_docx(case: dict, out_path: Path) -> None:
    doc = Document()
    footnotes: list[str] = []
    for block in case["blocks"]:
        kind = block["kind"]
        if kind == "title":
            doc.add_heading(block["text"], level=0)
        elif kind == "heading":
            doc.add_heading(block["text"], level=block.get("level", 1))
        elif kind == "paragraph":
            doc.add_paragraph(block["text"])
        elif kind == "list_item":
            doc.add_paragraph(block["text"], style=block.get("style", "List Bullet"))
        elif kind == "table":
            _build_docx_table(doc, block)
        elif kind == "caption":
            doc.add_paragraph(block["text"], style="Caption")
        elif kind == "image":
            doc.add_picture(_png_stream(), width=Inches(block.get("width", 2.0)))
        elif kind == "footnote":
            footnotes.append(block["text"])
        elif kind == "header":
            sec = doc.sections[0]
            (sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()).text = block["text"]
        elif kind == "footer":
            sec = doc.sections[0]
            (sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()).text = block["text"]
        elif kind == "page_break":
            doc.add_page_break()
        elif kind == "textbox":
            _add_textbox(doc, block["text"])
        elif kind == "bad_heading":
            p = doc.add_paragraph()
            run = p.add_run(block.get("number", "") + (" " if block.get("number") else "") + block["text"])
            run.bold = block.get("bold", True)
            if block.get("size"):
                run.font.size = Pt(block["size"])
        else:
            doc.add_paragraph(block.get("text", ""))
    doc.save(str(out_path))
    if footnotes:
        _inject_footnotes(out_path, footnotes)


def _build_docx_table(doc: Document, block: dict) -> None:
    headers = block.get("headers", [])
    rows = block.get("rows", [])
    n_cols = len(headers[0]) if headers else (len(rows[0]) if rows else 1)
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"
    for hrow in headers:
        cells = table.add_row().cells
        for i, val in enumerate(hrow):
            cells[i].text = str(val)
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val) if val is not None else ""
    for merge in block.get("merges", []):
        a, b = merge
        # 解析 A1 形式的左上/右下
        def _col_letter(ref: str) -> int:
            s = "".join(ch for ch in ref if ch.isalpha())
            n = 0
            for ch in s:
                n = n * 26 + (ord(ch.upper()) - 64)
            return n - 1
        def _row_num(ref: str) -> int:
            return int("".join(ch for ch in ref if ch.isdigit())) - 1
        ca, ra = _col_letter(a), _row_num(a)
        cb, rb = _col_letter(b), _row_num(b)
        table.cell(ra, ca).merge(table.cell(rb, cb))


def _inject_footnotes(path: Path, footnotes: list[str]) -> None:
    """把脚注写进 word/footnotes.xml 并建立 document.xml → footnotes.xml 关系。

    python-docx 1.2.0 无公开 footnote API，故保存后直接操作 zip 包，写入标准
    footnotes.xml（含 separator / continuationSeparator / 真实脚注）。解析器与
    selfcheck 直接读 footnotes.xml，因此 14A 脚注链路闭合。无须在正文插入
    footnoteReference（基准只需脚注文本可被抽取/自证）。
    """
    REL_FOOTNOTES = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
    )
    R_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}

    # 构建 footnotes.xml（separator / continuationSeparator 必须有）
    fn_items = [
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>',
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        '<w:continuationSeparator/></w:r></w:p></w:footnote>',
    ]
    for i, text in enumerate(footnotes, start=1):
        fid = str(2 + i)  # 3,4,...
        safe = _xml_escape(text)
        fn_items.append(
            f'<w:footnote w:id="{fid}"><w:p><w:r>'
            f'<w:t xml:space="preserve">{safe}</w:t></w:r></w:p></w:footnote>'
        )
    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        + "".join(fn_items)
        + "</w:footnotes>"
    )
    data["word/footnotes.xml"] = footnotes_xml.encode("utf-8")

    # 在 document.xml.rels 中建立指向 footnotes.xml 的关系（找最小未用 rId）
    rels_xml = ET.fromstring(data["word/_rels/document.xml.rels"])
    used = {rel.get("Id") for rel in rels_xml.findall(f"{{{R_NS}}}Relationship")}
    n = 1
    while f"rId{n}" in used:
        n += 1
    el = ET.SubElement(rels_xml, f"{{{R_NS}}}Relationship")
    el.set("Id", f"rId{n}")
    el.set("Type", REL_FOOTNOTES)
    el.set("Target", "footnotes.xml")
    ET.register_namespace("", R_NS)
    data["word/_rels/document.xml.rels"] = ET.tostring(rels_xml, encoding="UTF-8")

    # 在 [Content_Types].xml 中声明 footnotes 部件（否则严格读取器会拒绝）
    CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
    ct = ET.fromstring(data["[Content_Types].xml"])
    need_ct = True
    for ov in ct.findall(f"{{{CT_NS}}}Override"):
        if ov.get("PartName") == "/word/footnotes.xml":
            need_ct = False
            break
    if need_ct:
        ov = ET.SubElement(ct, f"{{{CT_NS}}}Override")
        ov.set("PartName", "/word/footnotes.xml")
        ov.set(
            "ContentType",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        )
    ET.register_namespace("", CT_NS)
    data["[Content_Types].xml"] = ET.tostring(ct, encoding="UTF-8")

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in data:  # 必须遍历 data（含新写入的 footnotes.xml）
            zout.writestr(name, data[name])
    tmp.replace(path)


# --------------------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------------------


def _inject_cached_values(path: Path, mapping: dict[str, dict[str, object]]) -> None:
    """把 formula 单元格补上 cached <v>，使文件同时含 <f> 与 <v>（openpyxl 自身做不到）。"""

    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    # 建立 sheet name -> sheetN.xml 映射
    wb_xml = ET.fromstring(data["xl/workbook.xml"])
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    sheet_tags = wb_xml.findall(".//a:sheet", ns)
    name_to_file = {}
    for i, st in enumerate(sheet_tags, start=1):
        sname = st.get("name")
        name_to_file[sname] = f"xl/worksheets/sheet{i}.xml"
    for sname, cells in mapping.items():
        fn = name_to_file.get(sname)
        if not fn or fn not in data:
            continue
        root = ET.fromstring(data[fn])
        for c in root.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
            ref = c.get("r")
            if ref in cells:
                f = c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f")
                if f is not None:
                    v_el = c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v")
                    if v_el is None:
                        v_el = ET.SubElement(c, "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v")
                    # 即使 openpyxl 已写出空的 <v/>，也要补上缓存值（双表示之一）
                    if not v_el.text:
                        v_el.text = str(cells[ref])
        data[fn] = ET.tostring(root, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    tmp.replace(path)


def _build_xlsx(case: dict, out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule

    wb = Workbook()
    wb.remove(wb.active)
    formula_injections: dict[str, dict[str, object]] = {}

    for sheet in case["sheets"]:
        ws = wb.create_sheet(title=sheet["name"])
        if sheet.get("state") == "hidden":
            ws.sheet_state = "hidden"
        # notes
        for note in sheet.get("notes", []):
            ws[note["cell"]] = note["text"]
        for table in sheet.get("tables", []):
            _fill_table(ws, table)
        for note in sheet.get("side_notes", []):
            ws[note["cell"]] = note["text"]
        # hidden rows/cols
        for r in sheet.get("hidden_rows", []):
            ws.row_dimensions[r].hidden = True
        for c in sheet.get("hidden_cols", []):
            ws.column_dimensions[c].hidden = True
        # named table
        for table in sheet.get("tables", []):
            if table.get("is_named"):
                ref = _table_range(ws, table)
                tbl = Table(displayName=table.get("table_name", "Table1"), ref=ref)
                tbl.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2", showRowStripes=True
                )
                ws.add_table(tbl)
        # formula injections
        for table in sheet.get("tables", []):
            fc = table.get("formula_cells")
            if fc:
                formula_injections.setdefault(sheet["name"], {}).update(
                    {k: v[1] for k, v in fc.items()}
                )
        # chart
        ch = sheet.get("chart")
        if ch:
            data_ref = None
            for table in sheet.get("tables", []):
                data_ref = table
            if data_ref:
                anchor_cell = ch.get("anchor", "D1")
                cat = Reference(ws, min_col=_col(ch["cat_col"]), min_row=2,
                                max_row=1 + len(data_ref.get("rows", [])))
                val = Reference(ws, min_col=_col(ch["val_col"]), min_row=1,
                                max_row=1 + len(data_ref.get("rows", [])))
                chart = BarChart()
                chart.title = ch.get("title", "chart")
                chart.add_data(val, titles_from_data=True)
                chart.set_categories(cat)
                ws.add_chart(chart, anchor_cell)
        # conditional formatting
        cf = sheet.get("conditional_format")
        if cf:
            ws.conditional_formatting.add(
                cf["range"], ColorScaleRule(start_type="min", end_type="max")
            )

    wb.save(str(out_path))
    if formula_injections:
        _inject_cached_values(out_path, formula_injections)


def _col(letter: str) -> int:
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n


def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _table_range(ws, table: dict) -> str:
    anchor = table["anchor"]
    ac = _col("".join(ch for ch in anchor if ch.isalpha()))
    ar = int("".join(ch for ch in anchor if ch.isdigit()))
    # headers 为扁平列名列表时只有 1 个表头行；多级表头为 2 行
    n_cols = len(table["headers"]) if not table.get("multi_level_header") else len(table["header_bottom"])
    n_header_rows = 2 if table.get("multi_level_header") else 1
    n_rows = n_header_rows + len(table.get("rows", []))
    return f"{anchor}:{_col_letter(ac + n_cols - 1)}{ar + n_rows - 1}"


def _fill_table(ws, table: dict) -> None:
    anchor = table["anchor"]
    ac = _col("".join(ch for ch in anchor if ch.isalpha()))
    ar = int("".join(ch for ch in anchor if ch.isdigit()))
    r = ar
    if table.get("multi_level_header"):
        for ci, val in enumerate(table["header_top"]):
            ws.cell(row=r, column=ac + ci, value=val)
        r += 1
        for ci, val in enumerate(table["header_bottom"]):
            ws.cell(row=r, column=ac + ci, value=val)
        r += 1
    else:
        for ci, val in enumerate(table["headers"]):
            ws.cell(row=r, column=ac + ci, value=val)
        r += 1
    for row in table.get("rows", []):
        for ci, val in enumerate(row):
            ws.cell(row=r, column=ac + ci, value=val)
        r += 1
    # merges
    for merge in table.get("merges", []):
        ws.merge_cells(f"{merge[0]}:{merge[1]}")
    # formulas
    for cell_ref, (formula, _cached) in (table.get("formula_cells") or {}).items():
        ws[cell_ref] = formula


def main() -> int:
    out_dir = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else (
        sandbox14a.ARTIFACT_ROOT / "controlled_corpus"
    )
    docx_dir = out_dir / "docx"
    xlsx_dir = out_dir / "xlsx"
    sandbox14a.ensure_dir(docx_dir)
    sandbox14a.ensure_dir(xlsx_dir)

    built = {"docx": [], "xlsx": []}
    for case in corpus14a.DOCX_CASES:
        p = docx_dir / f"{case['case_id']}.docx"
        _build_docx(case, p)
        built["docx"].append(p.name)
    for case in corpus14a.XLSX_CASES:
        p = xlsx_dir / f"{case['case_id']}.xlsx"
        _build_xlsx(case, p)
        built["xlsx"].append(p.name)

    summary = {
        "artifact_type": "phase14a_corpus_build",
        "docx_count": len(built["docx"]),
        "xlsx_count": len(built["xlsx"]),
        "docx": built["docx"],
        "xlsx": built["xlsx"],
    }
    sandbox14a.save_json(summary, out_dir / "build_summary.json")
    print(sandbox14a.sha256_json(summary))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
