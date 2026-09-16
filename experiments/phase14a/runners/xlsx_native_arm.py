"""Phase 14A-A 原生 baseline：XLSX (X-A0) openpyxl native。

把 .xlsx 读成统一 StructuredDocument（schema14a）。重点保留（方案 §2.5 / X-G*）：
- **formula + cached_value 双表示**（data_only=False 取公式，data_only=True 取缓存值）
- 每个非空 cell 的 coordinate / value / type 不得静默丢失（X-G1）
- merged ranges / named table / 启发式子表边界（X-G3/X-G4）
- sheet name/order/hidden（X-G5）、hidden row/col（X-G6）、display vs raw（number_format，X-G7）
- 不伪造 page / bbox（X-G1 反向）

运行：python -m phase14a.runners.xlsx_native_arm --xlsx <f> --out <f.json>
"""

from __future__ import annotations

import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from phase14a import sandbox14a, schema14a


def _detect_tables(ws):
    """返回 [(min_row,max_row,min_col,max_col,inferred,table_name)]。"""
    nonempty = [
        (c.row, c.column)
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    ]
    if not nonempty:
        return []
    rows_present = sorted({r for r, _ in nonempty})
    groups: list[list[int]] = []
    cur = [rows_present[0]]
    for r in rows_present[1:]:
        if r - cur[-1] > 1:
            groups.append(cur)
            cur = [r]
        else:
            cur.append(r)
    groups.append(cur)

    named = {}
    for name, tbl in ws.tables.items():
        # openpyxl 版本差异：表格值可能是 Table 对象（取 .ref）或直接的 ref 字符串
        ref = tbl.ref if hasattr(tbl, "ref") else tbl
        a, b = ref.split(":")
        ra = int("".join(ch for ch in a if ch.isdigit()))
        rb = int("".join(ch for ch in b if ch.isdigit()))
        ca = ord("".join(ch for ch in a if ch.isalpha()).upper()) - 64
        cb = ord("".join(ch for ch in b if ch.isalpha()).upper()) - 64
        named[(ra, rb, ca, cb)] = name

    result = []
    for g in groups:
        rmin, rmax = g[0], g[-1]
        cols = sorted({c for r, c in nonempty if rmin <= r <= rmax})
        cmin, cmax = cols[0], cols[-1]
        key = (rmin, rmax, cmin, cmax)
        if key in named:
            result.append((*key, False, named[key]))
        else:
            result.append((*key, True, None))
    return result


def parse(xlsx_path: Path) -> dict:
    # 受控基准文件很小，关闭 read_only 以保留 merged_cells / tables / row_dimensions 等结构 API。
    wb_f = load_workbook(str(xlsx_path), data_only=False)
    wb_v = load_workbook(str(xlsx_path), data_only=True)
    doc_id = xlsx_path.stem
    elements: list[dict] = []

    elements.append(schema14a.make_element(
        document_id=doc_id, element_id=f"el-xlsx-wb",
        type="workbook_metadata", raw_text=f"workbook:{doc_id}",
        source_format="xlsx",
        source_ref={"sheet": None, "range": None, "cell": None, "table_name": None, "inferred": None},
        source_parser="X-A0_openpyxl", retrieval_role="workbook_metadata",
        metadata={"sheet_count": len(wb_f.sheetnames), "sheet_names": list(wb_f.sheetnames)},
    ))

    for order, sheet_name in enumerate(wb_f.sheetnames):
        ws_f = wb_f[sheet_name]
        ws_v = wb_v[sheet_name]
        state = ws_f.sheet_state
        hidden_flag = state == "hidden"
        merged = [str(m) for m in ws_f.merged_cells.ranges]

        elements.append(schema14a.make_element(
            document_id=doc_id, element_id=f"el-xlsx-sheet-{order}",
            type="sheet", raw_text=f"sheet:{sheet_name}",
            source_format="xlsx",
            source_ref={"sheet": sheet_name, "range": None, "cell": None, "table_name": None, "inferred": None},
            source_parser="X-A0_openpyxl", retrieval_role="sheet",
            metadata={"name": sheet_name, "order": order, "state": state,
                      "hidden": hidden_flag, "merged_ranges": merged},
        ))

        tables = _detect_tables(ws_f)
        table_by_cell: dict[tuple[int, int], dict] = {}
        for (rmin, rmax, cmin, cmax, inferred, tname) in tables:
            rng = f"{get_column_letter(cmin)}{rmin}:{get_column_letter(cmax)}{rmax}"
            elements.append(schema14a.make_element(
                document_id=doc_id, element_id=f"el-xlsx-tbl-{order}-{rmin}",
                type="table", raw_text=f"table:{rng}",
                source_format="xlsx",
                source_ref={"sheet": sheet_name, "range": rng, "cell": None,
                            "table_name": tname, "inferred": inferred},
                source_parser="X-A0_openpyxl", retrieval_role="table",
                metadata={"rows": rmax - rmin + 1, "cols": cmax - cmin + 1,
                          "inferred": inferred, "named": tname is not None},
            ))
            for rr in range(rmin, rmax + 1):
                for cc in range(cmin, cmax + 1):
                    table_by_cell[(rr, cc)] = {"range": rng, "inferred": inferred}

        # 单元格层
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                coord = cell.coordinate
                r, c = cell.row, cell.column
                is_formula = cell.data_type == "f"
                formula = cell.value if is_formula else None
                cached = ws_v[coord].value if not is_formula else (ws_v[coord].value if ws_v[coord].value is not None else None)
                # 对于公式单元格，data_only=True 给出缓存值
                if is_formula:
                    cached = ws_v[coord].value
                tbl_ref = table_by_cell.get((r, c), {}).get("range")
                tmeta = table_by_cell.get((r, c), {})
                hidden_row = ws_f.row_dimensions[r].hidden if r in ws_f.row_dimensions else False
                hidden_col = ws_f.column_dimensions[get_column_letter(c)].hidden if get_column_letter(c) in ws_f.column_dimensions else False
                elements.append(schema14a.make_element(
                    document_id=doc_id, element_id=f"el-xlsx-cell-{order}-{coord}",
                    type="cell", raw_text=str(cell.value),
                    source_format="xlsx",
                    source_ref={"sheet": sheet_name, "range": tbl_ref, "cell": coord,
                                "table_name": tmeta.get("table_name"), "inferred": tmeta.get("inferred")},
                    source_parser="X-A0_openpyxl", retrieval_role="cell",
                    metadata={
                        "value": formula if is_formula else cell.value,
                        "formula": formula,
                        "cached_value": cached,
                        "data_type": cell.data_type,
                        "number_format": cell.number_format,
                        "is_formula": is_formula,
                        "is_hidden_row": bool(hidden_row),
                        "is_hidden_col": bool(hidden_col),
                        "in_table": tbl_ref is not None,
                    },
                ))

        wb_f.close()
        wb_v.close()

    return schema14a.make_document(
        document_id=doc_id, source_format="xlsx", source_file=xlsx_path.name,
        elements=elements, document_metadata={"parser": "X-A0_openpyxl"},
        source_sha256=sandbox14a.sha256_file(xlsx_path),
    )


def main() -> int:
    xlsx_path = Path(sys.argv[sys.argv.index("--xlsx") + 1])
    out_path = Path(sys.argv[sys.argv.index("--out") + 1])
    doc = parse(xlsx_path)
    sandbox14a.save_json(doc, out_path)
    print(f"X-A0 parsed {xlsx_path.name}: {doc['element_count']} elements")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
