"""Phase 14A-A 原生 baseline：DOCX (W-A0) python-docx native。

把 .docx 读成统一 StructuredDocument（schema14a）。重点保留：
- 显式 OOXML style → semantic role（方案 §2.2 / W-G4）
- heading hierarchy / paragraph order / list / table 结构
- header/footer 分离（W-G5）、footnote、image/caption、textbox
- 不强求伪造 page+bbox（W-G7）

运行：python -m phase14a.runners.docx_native_arm --docx <f> --out <f.json>
"""

from __future__ import annotations

import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.document import Document as _Doc
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from phase14a import sandbox14a, schema14a

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _classify(para: Paragraph) -> str:
    style = (para.style.name if para.style else "Normal") or "Normal"
    if style == "Title":
        return "document_title"
    if style.startswith("Heading"):
        return "heading"
    if style in ("List Bullet", "List Number", "List Bullet 2", "List Number 2",
                 "List Bullet 3", "List Number 3"):
        return "list_item"
    if style == "Caption":
        return "caption"
    if style == "FootnoteText":
        return "footnote"
    return "paragraph"


def _read_footnotes(docx_path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(docx_path) as z:
            if "word/footnotes.xml" not in z.namelist():
                return []
            root = ET.fromstring(z.read("word/footnotes.xml"))
    except Exception:
        return []
    texts: list[str] = []
    for fn in root.iter(f"{{{W_NS}}}footnote"):
        if fn.get(f"{{{W_NS}}}type"):  # 跳过 separator/continuation
            continue
        texts.append("".join(t.text or "" for t in fn.iter(f"{{{W_NS}}}t")))
    return [t for t in texts if t]


def _read_textboxes(docx_path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(docx_path) as z:
            root = ET.fromstring(z.read("word/document.xml"))
    except Exception:
        return []
    boxes: list[str] = []
    for txbx in root.iter(f"{{{W_NS}}}txbxContent"):
        boxes.append("".join(t.text or "" for t in txbx.iter(f"{{{W_NS}}}t")))
    return boxes


def parse(docx_path: Path) -> dict:
    doc = Document(str(docx_path))
    elements: list[dict] = []
    doc_id = docx_path.stem
    pid = 0
    table_idx = -1

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            style = (block.style.name if block.style else "Normal") or "Normal"
            etype = _classify(block)
            has_pic = "a:blip" in block._p.xml or "pic:pic" in block._p.xml
            sr = {
                "paragraph_index": pid,
                "xml_path": f"/w:document/w:body/w:p[{pid + 1}]",
                "style_id": style,
                "table_index": None,
                "cell": None,
            }
            metadata = {}
            if has_pic:
                # 同时记录 image 元素，但保持段落顺序
                elements.append(schema14a.make_element(
                    document_id=doc_id,
                    element_id=f"el-docx-img-{pid}",
                    type="image",
                    raw_text="",
                    source_format="docx",
                    source_ref=sr,
                    source_parser="W-A0_python_docx",
                    retrieval_role="image",
                    metadata={"near_paragraph_index": pid},
                ))
            elements.append(schema14a.make_element(
                document_id=doc_id,
                element_id=f"el-docx-p-{pid}",
                type=etype,
                raw_text=block.text or "",
                source_format="docx",
                source_ref=sr,
                source_parser="W-A0_python_docx",
                retrieval_role=("furniture" if etype in ("header", "footer") else "body"),
                metadata=metadata,
            ))
            pid += 1
        else:  # Table
            table_idx += 1
            rows = block.rows
            n_rows = len(rows)
            n_cols = len(rows[0].cells) if n_rows else 0
            elements.append(schema14a.make_element(
                document_id=doc_id,
                element_id=f"el-docx-tbl-{table_idx}",
                type="table",
                raw_text="",
                source_format="docx",
                source_ref={"paragraph_index": None, "xml_path": f"table[{table_idx + 1}]",
                            "style_id": None, "table_index": table_idx, "cell": None},
                source_parser="W-A0_python_docx",
                retrieval_role="table",
                metadata={"rows": n_rows, "cols": n_cols},
            ))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row.cells):
                    elements.append(schema14a.make_element(
                        document_id=doc_id,
                        element_id=f"el-docx-tbl{table_idx}-c{r}-{c}",
                        type="table_cell",
                        raw_text=cell.text or "",
                        source_format="docx",
                        source_ref={"paragraph_index": None, "xml_path": None,
                                    "style_id": None, "table_index": table_idx,
                                    "cell": f"R{r + 1}C{c + 1}"},
                        source_parser="W-A0_python_docx",
                        retrieval_role="table_cell",
                        metadata={},
                    ))

    # header / footer
    for si, section in enumerate(doc.sections):
        for hp in section.header.paragraphs:
            if hp.text.strip():
                elements.append(schema14a.make_element(
                    document_id=doc_id, element_id=f"el-docx-h-{si}",
                    type="header", raw_text=hp.text, source_format="docx",
                    source_ref={"paragraph_index": None, "xml_path": "header",
                                "style_id": None, "table_index": None, "cell": None},
                    source_parser="W-A0_python_docx", retrieval_role="furniture",
                    metadata={}))
        for fp in section.footer.paragraphs:
            if fp.text.strip():
                elements.append(schema14a.make_element(
                    document_id=doc_id, element_id=f"el-docx-f-{si}",
                    type="footer", raw_text=fp.text, source_format="docx",
                    source_ref={"paragraph_index": None, "xml_path": "footer",
                                "style_id": None, "table_index": None, "cell": None},
                    source_parser="W-A0_python_docx", retrieval_role="furniture",
                    metadata={}))

    # footnotes
    for fi, ftext in enumerate(_read_footnotes(docx_path)):
        elements.append(schema14a.make_element(
            document_id=doc_id, element_id=f"el-docx-fn-{fi}",
            type="footnote", raw_text=ftext, source_format="docx",
            source_ref={"paragraph_index": None, "xml_path": "footnotes",
                        "style_id": None, "table_index": None, "cell": None},
            source_parser="W-A0_python_docx", retrieval_role="footnote", metadata={}))

    # textboxes
    for bi, btext in enumerate(_read_textboxes(docx_path)):
        elements.append(schema14a.make_element(
            document_id=doc_id, element_id=f"el-docx-tb-{bi}",
            type="paragraph", raw_text=btext, source_format="docx",
            source_ref={"paragraph_index": None, "xml_path": "textbox",
                        "style_id": "Textbox", "table_index": None, "cell": None},
            source_parser="W-A0_python_docx", retrieval_role="body",
            metadata={"is_textbox": True}))

    return schema14a.make_document(
        document_id=doc_id, source_format="docx", source_file=docx_path.name,
        elements=elements, document_metadata={"parser": "W-A0_python_docx"},
        source_sha256=sandbox14a.sha256_file(docx_path),
    )


def main() -> int:
    docx_path = Path(sys.argv[sys.argv.index("--docx") + 1])
    out_path = Path(sys.argv[sys.argv.index("--out") + 1])
    doc = parse(docx_path)
    sandbox14a.save_json(doc, out_path)
    print(f"W-A0 parsed {docx_path.name}: {doc['element_count']} elements")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
